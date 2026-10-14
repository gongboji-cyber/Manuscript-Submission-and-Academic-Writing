#!/usr/bin/env python3
"""Generate software-copyright DOCX materials from confirmed project facts.

Inputs:
- canonical application facts JSON
- confirmed manual Markdown
- confirmed first-party code selection JSON
- optional screenshot manifest CSV

Outputs:
- <software>_软件说明书.docx
- <software>_源程序鉴别材料.docx
- DOCX_BUILD_MANIFEST.json

The script never invents source code or legal facts. Selected source files must already
be marked FIRST_PARTY/APPLICANT_CONFIRMED/VERIFIED_FIRST_PARTY.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:
    raise SystemExit("python-docx is required: pip install python-docx") from exc

ALLOWED_OWNERSHIP={"FIRST_PARTY","APPLICANT_CONFIRMED","VERIFIED_FIRST_PARTY"}
PLACEHOLDER_RE=re.compile(r"(?:TODO|TBD|待确认|提交前更新|截图待补|placeholder)",re.I)
IMAGE_TOKEN_RE=re.compile(r"^\[\[FIGURE:([^\]]+)\]\]\s*$")
MD_IMAGE_RE=re.compile(r"^!\[[^\]]*\]\(([^)]+)\)\s*$")

def sha256(path:Path)->str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024),b""): h.update(chunk)
    return h.hexdigest()

def set_cell_margins(cell,top=80,start=80,bottom=80,end=80):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{m}"))
        if node is None:
            node=OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=paragraph.add_run()
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE")
    run._r.addnext(fld)

def configure_document(doc:Document,title:str,code:bool=False):
    sec=doc.sections[0]
    sec.top_margin=Cm(2.0 if code else 2.5); sec.bottom_margin=Cm(2.0 if code else 2.5)
    sec.left_margin=Cm(2.0 if code else 2.5); sec.right_margin=Cm(2.0 if code else 2.5)
    sec.header_distance=Cm(0.8); sec.footer_distance=Cm(0.8)
    hp=sec.header.paragraphs[0]; hp.text=title; hp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.name="宋体"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体"); r.font.size=Pt(9)
    add_page_number(sec.footer.paragraphs[0])

    styles=doc.styles
    normal=styles["Normal"]; normal.font.name="宋体"; normal._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
    normal.font.size=Pt(12 if not code else 7.5); normal.paragraph_format.space_after=Pt(0 if code else 3); normal.paragraph_format.line_spacing=1.0 if code else 1.5
    for sname,size in (("Title",18),("Heading 1",14),("Heading 2",12),("Heading 3",12)):
        st=styles[sname]; st.font.name="黑体"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"黑体"); st.font.size=Pt(size)

def read_facts(path:Path)->dict:
    obj=json.loads(path.read_text(encoding="utf-8")); c=obj.get("canonical",obj)
    missing=[k for k in ("software_full_name","version") if not str(c.get(k,"")).strip()]
    if missing: raise SystemExit("missing canonical facts: "+", ".join(missing))
    for k in ("software_full_name","version","copyright_owner","completion_date"):
        if PLACEHOLDER_RE.search(str(c.get(k,""))): raise SystemExit(f"canonical fact still unresolved: {k}={c.get(k)!r}")
    return c

def load_screenshots(path:Path|None)->dict[str,Path]:
    if not path: return {}
    rows=list(csv.DictReader(path.open(encoding="utf-8-sig",newline=""))); out={}
    for r in rows:
        key=(r.get("figure_id") or r.get("id") or "").strip(); raw=(r.get("filename") or r.get("path") or "").strip()
        if key and raw:
            p=Path(raw)
            if not p.is_absolute(): p=path.parent/p
            if p.is_file(): out[key]=p.resolve()
    return out

def add_inline_markdown(paragraph,text:str):
    token=re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)"); pos=0
    for m in token.finditer(text):
        if m.start()>pos: paragraph.add_run(text[pos:m.start()])
        piece=m.group(0)
        if piece.startswith("`"):
            r=paragraph.add_run(piece[1:-1]); r.font.name="Consolas"; r.font.size=Pt(10)
        else:
            r=paragraph.add_run(piece[2:-2]); r.bold=True
        pos=m.end()
    if pos<len(text): paragraph.add_run(text[pos:])

def iter_markdown_blocks(lines:list[str]):
    i=0
    while i<len(lines):
        raw=lines[i].rstrip("\n")
        if not raw.strip(): i+=1; continue
        if raw.startswith("```"):
            code=[]; i+=1
            while i<len(lines) and not lines[i].startswith("```"):
                code.append(lines[i].rstrip("\n")); i+=1
            i+=1; yield ("code",code); continue
        if raw.startswith("#"):
            m=re.match(r"^(#{1,6})\s+(.*)$",raw)
            if m: yield ("heading",len(m.group(1)),m.group(2).strip()); i+=1; continue
        im=IMAGE_TOKEN_RE.match(raw.strip())
        if im: yield ("image",im.group(1).strip()); i+=1; continue
        mm=MD_IMAGE_RE.match(raw.strip())
        if mm: yield ("image_path",mm.group(1).strip()); i+=1; continue
        if raw.lstrip().startswith(("- ","* ")):
            items=[]
            while i<len(lines) and lines[i].lstrip().startswith(("- ","* ")):
                items.append(lines[i].lstrip()[2:].rstrip()); i+=1
            yield ("bullets",items); continue
        if "|" in raw and i+1<len(lines) and re.match(r"^\s*\|?\s*:?-+",lines[i+1]):
            table=[raw]; i+=2
            while i<len(lines) and "|" in lines[i] and lines[i].strip():
                table.append(lines[i].rstrip()); i+=1
            yield ("table",table); continue
        para=[raw.strip()]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith("#") and not lines[i].startswith("```") and not IMAGE_TOKEN_RE.match(lines[i].strip()) and not lines[i].lstrip().startswith(("- ","* ")):
            para.append(lines[i].strip()); i+=1
        yield ("paragraph"," ".join(para))

def build_manual(facts:dict,manual_md:Path,screenshot_manifest:Path|None,out:Path)->dict:
    text=manual_md.read_text(encoding="utf-8")
    if PLACEHOLDER_RE.search(text): raise SystemExit("manual markdown contains unresolved placeholder/TODO")
    shots=load_screenshots(screenshot_manifest); doc=Document(); title=f"{facts['software_full_name']} {facts['version']}"
    configure_document(doc,title,code=False)

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title); r.bold=True; r.font.size=Pt(20); r.font.name="黑体"; r._element.rPr.rFonts.set(qn("w:eastAsia"),"黑体")
    p2=doc.add_paragraph(); p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
    rr=p2.add_run("软件使用说明书"); rr.bold=True; rr.font.size=Pt(16)
    doc.add_paragraph("")
    for k,v in (("软件全称",facts.get("software_full_name")),("软件简称",facts.get("software_short_name")),("版本号",facts.get("version")),("著作权人",facts.get("copyright_owner")),("开发完成日期",facts.get("completion_date"))):
        if v:
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(f"{k}：{v}")
    doc.add_page_break()

    for block in iter_markdown_blocks(text.splitlines()):
        kind=block[0]
        if kind=="heading":
            _,level,heading=block; doc.add_heading(heading,level=min(level,3))
        elif kind=="paragraph":
            p=doc.add_paragraph(); add_inline_markdown(p,block[1])
        elif kind=="bullets":
            for item in block[1]:
                p=doc.add_paragraph(style="List Bullet"); add_inline_markdown(p,item)
        elif kind=="code":
            for line in block[1]:
                p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run(line); r.font.name="Consolas"; r.font.size=Pt(9)
        elif kind=="image":
            fig=block[1]
            if fig not in shots: raise SystemExit(f"screenshot missing for figure id: {fig}")
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(shots[fig]),width=Cm(15.5))
        elif kind=="image_path":
            pth=Path(block[1])
            if not pth.is_absolute(): pth=manual_md.parent/pth
            if not pth.is_file(): raise SystemExit(f"markdown image not found: {pth}")
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(pth),width=Cm(15.5))
        elif kind=="table":
            rows=[[c.strip() for c in line.strip().strip("|").split("|")] for line in block[1]]; cols=max(len(r) for r in rows)
            t=doc.add_table(rows=len(rows),cols=cols); t.style="Table Grid"
            for i,row in enumerate(rows):
                for j,val in enumerate(row):
                    t.cell(i,j).text=val; set_cell_margins(t.cell(i,j))
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
    return {"path":str(out),"sha256":sha256(out),"screenshots":{k:str(v) for k,v in shots.items()}}

def selected_source_files(project:Path,selection:Path)->list[tuple[str,Path]]:
    data=json.loads(selection.read_text(encoding="utf-8")); files=[]
    for rec in data.get("files",[]):
        if not rec.get("selected"): continue
        if rec.get("ownership_status") not in ALLOWED_OWNERSHIP: raise SystemExit(f"selected file is not confirmed first-party: {rec.get('path')}")
        p=project/rec["path"]
        if not p.is_file(): raise SystemExit(f"selected source file missing: {p}")
        files.append((rec["path"],p))
    if not files: raise SystemExit("no confirmed first-party source files selected")
    return files

def build_source(facts:dict,project:Path,selection:Path,out:Path,lines_per_page:int)->dict:
    doc=Document(); title=f"{facts['software_full_name']} {facts['version']}"; configure_document(doc,title,code=True)
    global_line=0; page=1; mapping=[]; page_segments=[]; page_count=0

    def add_code_line(text:str):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; p.paragraph_format.keep_with_next=False
        r=p.add_run(text if text else " "); r.font.name="Consolas"; r.font.size=Pt(7.5)

    files=selected_source_files(project,selection)
    total_remaining=sum(len(p.read_text(encoding="utf-8",errors="replace").splitlines()) for _,p in files); emitted=0
    for rel,path in files:
        lines=path.read_text(encoding="utf-8",errors="replace").splitlines(); local=0
        while local<len(lines):
            remaining=lines_per_page-page_count; take=min(remaining,len(lines)-local); start=local+1; end=local+take
            for line in lines[local:end]:
                add_code_line(line); global_line+=1; emitted+=1; page_count+=1
            page_segments.append(f"{rel}:{start}-{end}"); local=end
            if page_count==lines_per_page:
                mapping.append({"page":page,"expected_code_lines":page_count,"segments":" | ".join(page_segments),"status":"FULL_PAGE"})
                if emitted<total_remaining:
                    doc.add_page_break(); page+=1
                page_count=0; page_segments=[]
    if page_count:
        mapping.append({"page":page,"expected_code_lines":page_count,"segments":" | ".join(page_segments),"status":"FINAL_PARTIAL_PAGE"})
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
    map_path=out.with_name(out.stem+"_SOURCE_PAGE_MAP.csv")
    with map_path.open("w",newline="",encoding="utf-8-sig") as f:
        w=csv.DictWriter(f,fieldnames=["page","expected_code_lines","segments","status"]); w.writeheader(); w.writerows(mapping)
    return {"path":str(out),"sha256":sha256(out),"source_lines":global_line,"expected_pages":len(mapping),"page_map":str(map_path),"page_map_sha256":sha256(map_path)}

def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--project",required=True); ap.add_argument("--facts",required=True); ap.add_argument("--manual-md",required=True); ap.add_argument("--selection",required=True)
    ap.add_argument("--screenshots"); ap.add_argument("--out-dir",required=True); ap.add_argument("--code-lines-per-page",type=int,default=50)
    args=ap.parse_args()
    project=Path(args.project).resolve(); out_dir=Path(args.out_dir).resolve(); facts_path=Path(args.facts).resolve(); manual_md=Path(args.manual_md).resolve(); selection=Path(args.selection).resolve()
    shots=Path(args.screenshots).resolve() if args.screenshots else None; facts=read_facts(facts_path); full=facts["software_full_name"]
    manual_out=out_dir/f"{full}_软件说明书.docx"; source_out=out_dir/f"{full}_源程序鉴别材料.docx"
    manual=build_manual(facts,manual_md,shots,manual_out); source=build_source(facts,project,selection,source_out,args.code_lines_per_page)
    manifest={"schema_version":1,"facts":str(facts_path),"manual_markdown":str(manual_md),"selection":str(selection),"manual":manual,"source":source,
              "warnings":["DOCX creation does not prove pagination. Run render_pagination_qa.py and visually inspect rendered pages before submission."]}
    mp=out_dir/"DOCX_BUILD_MANIFEST.json"; mp.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    print(manual_out); print(source_out); print(mp); return 0

if __name__=="__main__": raise SystemExit(main())
