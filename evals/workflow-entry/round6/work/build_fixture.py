from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "materials" / "input.docx"


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in edges.items():
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn("w:" + key), str(value))


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

title = doc.add_paragraph()
title.style = doc.styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Synthetic Table Formatting Check")
run.font.name = "Arial"
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph()
p.paragraph_format.first_line_indent = Inches(0.25)
p.paragraph_format.space_after = Pt(7)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("Synthetic report. Enrollment ")
bold = p.add_run("n=24")
bold.bold = True
p.add_run(". Do not edit.")

table = doc.add_table(rows=3, cols=3)
table.style = "Table Grid"
table.autofit = False
widths = [Inches(2.7), Inches(1.4), Inches(1.8)]
values = [
    ["Group", "Count", "Estimate (95% CI)"],
    ["Control", "12", "7.5 (5.0 to 10.0)"],
    ["Intervention", "12", "6.0 (4.0 to 8.0)"],
]
for row_idx, row in enumerate(table.rows):
    for col_idx, cell in enumerate(row.cells):
        cell.width = widths[col_idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = values[row_idx][col_idx]
        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        )
        shade(cell, "D9EAF7" if row_idx == 0 else "F2F2F2")
        set_cell_border(
            cell,
            top={"val": "single", "sz": "8", "color": "4472C4"},
            bottom={"val": "single", "sz": "8", "color": "4472C4"},
            left={"val": "single", "sz": "8", "color": "4472C4"},
            right={"val": "single", "sz": "8", "color": "4472C4"},
            insideH={"val": "single", "sz": "8", "color": "4472C4"},
            insideV={"val": "single", "sz": "8", "color": "4472C4"},
        )

after = doc.add_paragraph()
after.paragraph_format.space_before = Pt(9)
after.paragraph_format.keep_with_next = False
after.alignment = WD_ALIGN_PARAGRAPH.RIGHT
after.add_run("Reference marker: ")
italic = after.add_run("7.5%")
italic.italic = True
after.add_run("; preserve exactly.")

doc.add_section(WD_SECTION.CONTINUOUS)
doc.save(OUT)
print(OUT)

