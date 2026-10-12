from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


OUT = Path(__file__).resolve().parent.parent / "fixtures"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def set_all_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "777777")
        borders.append(el)


def base_doc(title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)
    p = doc.add_paragraph(title)
    p.style = doc.styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Synthetic public fixture for formatting validation. No research or patient data are included.")
    return doc


def normal_long_table() -> None:
    doc = base_doc("Long Table Pagination Fixture")
    caption = doc.add_paragraph("Table 1. Synthetic repeated-measures summary")
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=2, cols=4)
    table.autofit = False
    widths = [1.0, 2.9, 1.25, 1.35]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Record"
    table.cell(0, 1).text = "Measurement"
    table.cell(0, 2).merge(table.cell(0, 3)).text = "Estimate"
    table.cell(1, 1).text = "Variable"
    table.cell(1, 2).text = "Value"
    table.cell(1, 3).text = "95% interval"
    for cell in table.rows[0].cells + table.rows[1].cells:
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
    for i in range(1, 73):
        cells = table.add_row().cells
        cells[0].text = f"R{i:03d}"
        cells[1].text = f"Synthetic outcome label {i:03d} with a stable descriptive phrase"
        cells[2].text = f"{i / 17:.3f}"
        cells[3].text = f"[{i / 17 - 0.12:.3f}, {i / 17 + 0.12:.3f}]"
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_all_borders(table)
    doc.add_paragraph("Note. Values are deterministic synthetic text used only to verify pagination and preservation.")
    doc.add_paragraph("Sentinel after table: NORMAL-END-88421")
    doc.save(OUT / "normal-long-table.docx")


def overpage_row() -> None:
    doc = base_doc("Over Page Row Fixture")
    caption = doc.add_paragraph("Table 1. Synthetic narrative record")
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Narrative"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "E2F0D9")
    row = table.add_row()
    row.cells[0].text = "A001"
    narrative = row.cells[1]
    narrative.text = ""
    for i in range(1, 181):
        p = narrative.paragraphs[0] if i == 1 else narrative.add_paragraph()
        p.add_run(f"TOKEN-{i:03d}: This sentence must remain visible and extractable after formatting.")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
    set_all_borders(table)
    doc.add_paragraph("Note. The single body row intentionally exceeds one page.")
    doc.add_paragraph("Sentinel after table: OVERPAGE-END-33017")
    doc.save(OUT / "overpage-row.docx")


def layout_table() -> None:
    doc = base_doc("Layout Table Fixture")
    doc.add_paragraph("The table below is an author-information layout container, not a scientific data table.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Corresponding author"
    table.cell(0, 1).text = "Example Author"
    table.cell(1, 0).text = "Institution"
    table.cell(1, 1).text = "Example University"
    table.cell(2, 0).text = "Document role"
    table.cell(2, 1).text = "Title-page layout only"
    set_all_borders(table)
    doc.add_paragraph("Sentinel after layout: LAYOUT-END-55009")
    doc.save(OUT / "layout-table.docx")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    normal_long_table()
    overpage_row()
    layout_table()


if __name__ == "__main__":
    main()
