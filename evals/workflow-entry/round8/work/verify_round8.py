from __future__ import annotations

import hashlib
import json
import re
import subprocess
import argparse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parent.parent


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical(element) -> bytes:
    return etree.tostring(element, method="c14n", with_comments=True)


def external_blocks(doc: Document) -> list[str]:
    return [
        hashlib.sha256(canonical(child)).hexdigest()
        for child in doc._element.body.iterchildren()
        if child.tag != qn("w:tbl")
    ]


def physical_table_text(table) -> list[list[str]]:
    return [
        ["".join(tc.itertext()) for tc in tr.tc_lst]
        for tr in table._tbl.tr_lst
    ]


def merge_signature(table) -> list[list[object]]:
    out = []
    for r, tr in enumerate(table._tbl.tr_lst):
        for c, tc in enumerate(tr.tc_lst):
            tc_pr = tc.tcPr
            span = tc_pr.gridSpan.val if tc_pr is not None and tc_pr.gridSpan is not None else 1
            out.append([r, c, span, tc.vMerge])
    return out


def border_value(tc, edge: str) -> tuple[str | None, str | None]:
    borders = tc.tcPr.find(qn("w:tcBorders")) if tc.tcPr is not None else None
    value = borders.find(qn(f"w:{edge}")) if borders is not None else None
    if value is None:
        return None, None
    return value.get(qn("w:val")), value.get(qn("w:sz"))


def pdf_page_text(pdf: Path, page: int) -> str:
    proc = subprocess.run(
        ["pdftotext", "-f", str(page), "-l", str(page), "-layout", str(pdf), "-"],
        check=True,
        capture_output=True,
        text=True,
    )
    return proc.stdout


def canonical_page_count(directory: Path) -> int:
    return sum(1 for path in directory.glob("page-*.png") if re.fullmatch(r"page-\d+\.png", path.name))


def table_grid_widths(table) -> list[int | None]:
    return [
        int(col.get(qn("w:w"))) if col.get(qn("w:w")) is not None else None
        for col in table._tbl.tblGrid.gridCol_lst
    ]


def compare_docx(src_path: Path, out_path: Path) -> dict[str, object]:
    src = Document(src_path)
    out = Document(out_path)
    return {
        "source_sha256": sha256(src_path),
        "output_sha256": sha256(out_path),
        "paragraph_text_equal": [p.text for p in src.paragraphs] == [p.text for p in out.paragraphs],
        "external_blocks_equal": external_blocks(src) == external_blocks(out),
        "section_xml_equal": [canonical(s._sectPr) for s in src.sections] == [canonical(s._sectPr) for s in out.sections],
        "table_count_equal": len(src.tables) == len(out.tables),
        "table_dimensions_equal": [(len(t.rows), len(t.columns)) for t in src.tables] == [(len(t.rows), len(t.columns)) for t in out.tables],
        "table_text_equal": [physical_table_text(t) for t in src.tables] == [physical_table_text(t) for t in out.tables],
        "merge_signature_equal": [merge_signature(t) for t in src.tables] == [merge_signature(t) for t in out.tables],
    }


def bundle_check(case: str, variant: str) -> dict[str, object]:
    bundle = ROOT / "runs" / case / "bundle"
    manifest_path = ROOT / "bundles" / f"{case.split('-')[0]}-{variant}" / "bundle-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    actual = sorted(p.relative_to(bundle).as_posix() for p in bundle.rglob("*") if p.is_file())
    listed = sorted(item["bundle_path"] for item in manifest["files"])
    payload = "\n".join((bundle / rel).read_text(encoding="utf-8") for rel in actual if rel.endswith((".md", ".py")))
    entry_hashes = {
        rel: sha256(bundle / rel)
        for rel in actual
    }
    expected_hashes = {item["bundle_path"]: item["sha256"] for item in manifest["files"]}
    return {
        "actual_files": actual,
        "manifest_files": listed,
        "file_set_equal": actual == listed,
        "hashes_equal": entry_hashes == expected_hashes,
        "contains_checklist_md": any(Path(rel).name == "checklist.md" for rel in actual),
        "contains_F01_or_F20": "F01" in payload or "F20" in payload,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--variant", choices=("current", "revised"), default="current")
    args = parser.parse_args()
    variant = args.variant

    p01_run = ROOT / "runs" / f"p01-{variant}"
    p02_run = ROOT / "runs" / f"p02-{variant}"
    p03_run = ROOT / "runs" / f"p03-{variant}"
    p01_render = p01_run / ("render-qa" if variant == "current" else "qa-render")
    p02_render = p02_run / ("qa-render" if variant == "current" else "qa-formatted-v2")
    p01 = compare_docx(p01_run / "input.docx", p01_run / "formatted.docx")
    p02 = compare_docx(p02_run / "input.docx", p02_run / "formatted.docx")

    p01_doc = Document(p01_run / "formatted.docx")
    p01_rows = p01_doc.tables[0]._tbl.tr_lst
    p01.update({
        "header_flags": [
            row.trPr is not None and row.trPr.find(qn("w:tblHeader")) is not None
            for row in p01_rows[:4]
        ],
        "top_rules": [border_value(tc, "top") for tc in p01_rows[0].tc_lst],
        "header_bottom_rules": [border_value(tc, "bottom") for tc in p01_rows[1].tc_lst],
        "final_bottom_rules": [border_value(tc, "bottom") for tc in p01_rows[-1].tc_lst],
        "source_grid_widths": table_grid_widths(Document(p01_run / "input.docx").tables[0]),
        "output_grid_widths": table_grid_widths(p01_doc.tables[0]),
        "render_pages": canonical_page_count(p01_render),
        "pdf_headers_by_page": [
            all(token in pdf_page_text(p01_render / "formatted.pdf", page) for token in ("Record", "Variable", "Value", "95% interval"))
            for page in range(1, canonical_page_count(p01_render) + 1)
        ],
        "row_ids_by_page": [
            [token for token in (f"R{i:03d}" for i in range(1, 73)) if token in pdf_page_text(p01_render / "formatted.pdf", page)]
            for page in range(1, canonical_page_count(p01_render) + 1)
        ],
    })

    p02_doc = Document(p02_run / "formatted.docx")
    p02_rows = p02_doc.tables[0]._tbl.tr_lst
    p02_page_count = canonical_page_count(p02_render)
    pdf_text = "\n".join(pdf_page_text(p02_render / "formatted.pdf", page) for page in range(1, p02_page_count + 1))
    tokens = [f"TOKEN-{i:03d}" for i in range(1, 181)]
    p02.update({
        "header_has_tblHeader": p02_rows[0].trPr is not None and p02_rows[0].trPr.find(qn("w:tblHeader")) is not None,
        "header_has_cantSplit": p02_rows[0].trPr is not None and p02_rows[0].trPr.find(qn("w:cantSplit")) is not None,
        "body_has_cantSplit": p02_rows[1].trPr is not None and p02_rows[1].trPr.find(qn("w:cantSplit")) is not None,
        "source_grid_widths": table_grid_widths(Document(p02_run / "input.docx").tables[0]),
        "output_grid_widths": table_grid_widths(p02_doc.tables[0]),
        "render_pages": p02_page_count,
        "all_tokens_once": all(pdf_text.count(token) == 1 for token in tokens),
        "token_positions_in_order": [pdf_text.find(token) for token in tokens] == sorted(pdf_text.find(token) for token in tokens),
        "sentinel_visible": "OVERPAGE-END-33017" in pdf_text,
        "pdf_headers_by_page": [
            all(token in pdf_page_text(p02_render / "formatted.pdf", page) for token in ("Item", "Narrative"))
            for page in range(1, p02_page_count + 1)
        ],
    })

    result = {
        "runtime": {
            "python": subprocess.run([str(Path(__import__("sys").executable)), "--version"], capture_output=True, text=True).stdout.strip(),
            "python_docx": __import__("docx").__version__,
            "lxml": etree.LXML_VERSION,
            "libreoffice": subprocess.run(["soffice", "--version"], capture_output=True, text=True).stdout.strip(),
            "pdftoppm": subprocess.run(["pdftoppm", "-v"], capture_output=True, text=True).stderr.splitlines()[0],
        },
        "p01": p01,
        "p02": p02,
        "p03": {
            "input_sha256": sha256(p03_run / "input.docx"),
            "expected_input_sha256": "e7b80c318e311100a20a3314dc3fd92141ba41a44ba03aad575d15efb96269cb",
            "formatted_exists": (p03_run / "formatted.docx").exists(),
            "log_exists": (p03_run / "execution-log.md").exists(),
        },
        "bundles": {
            case: bundle_check(case, variant)
            for case in (f"p01-{variant}", f"p02-{variant}", f"p03-{variant}")
        },
    }
    output_name = "parent-verification.json" if variant == "current" else "parent-verification-revised.json"
    (ROOT / output_name).write_text(
        json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
