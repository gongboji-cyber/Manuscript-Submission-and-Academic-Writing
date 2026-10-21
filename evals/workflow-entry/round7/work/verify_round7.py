from __future__ import annotations

import hashlib
import ast
import json
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "materials" / "simple-table.docx"
BOUNDARY = ROOT / "materials" / "boundary-crossing-merge.docx"
SCRIPT = ROOT / "candidate" / "skills" / "Word文档格式" / "scripts" / "stage_task_bundle.py"
SKILL_SOURCE = SCRIPT.parents[1]
LOCAL_BUNDLE = ROOT / "bundles" / "local-e"
FULL_BUNDLE = ROOT / "bundles" / "full-c"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_manifest(root: Path, scope: str) -> dict:
    manifest = json.loads((root / "bundle-manifest.json").read_text(encoding="utf-8"))
    assert manifest["scope"] == scope
    actual = []
    for path in sorted(p for p in (root / "bundle").rglob("*") if p.is_file() and "__pycache__" not in p.parts):
        actual.append(path.relative_to(root / "bundle").as_posix())
    listed = [entry["bundle_path"] for entry in manifest["files"]]
    assert actual == listed
    for entry in manifest["files"]:
        path = root / "bundle" / entry["bundle_path"]
        assert digest(path) == entry["sha256"]
        assert path.stat().st_size == entry["bytes"]
    return {
        "scope": scope,
        "entries": manifest["files"],
        "manifest_sha256": digest(root / "bundle-manifest.json"),
    }


def verify_staging_imports() -> dict:
    tree = ast.parse(SCRIPT.read_text(encoding="utf-8"))
    modules = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            modules.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module and node.module != "__future__":
            modules.add(node.module.split(".")[0])
    allowed = {"argparse", "hashlib", "json", "os", "pathlib", "shutil", "tempfile"}
    assert modules <= allowed
    return {"imports": sorted(modules), "allowed_standard_library_modules": sorted(allowed), "standard_library_only": True}


def run_stage_tests() -> dict:
    python = "/opt/codex/runtimes/codex-primary-runtime/dependencies/python/bin/python3"
    with tempfile.TemporaryDirectory(prefix="round7-stage-") as temp:
        temp_root = Path(temp)
        local_one = temp_root / "local-one"
        local_two = temp_root / "local-two"
        full = temp_root / "full"
        commands = []
        for scope, output in (("local-three-line", local_one), ("local-three-line", local_two), ("full-audit", full)):
            command = [python, "-B", str(SCRIPT), "--source-root", str(SKILL_SOURCE), "--scope", scope, "--output", str(output)]
            result = subprocess.run(command, text=True, capture_output=True)
            commands.append({"argv": command, "returncode": result.returncode, "stdout": result.stdout, "stderr": result.stderr})
            assert result.returncode == 0
        assert (local_one / "bundle-manifest.json").read_bytes() == (local_two / "bundle-manifest.json").read_bytes()

        nonempty = temp_root / "nonempty"
        nonempty.mkdir()
        (nonempty / "sentinel.txt").write_text("preserve\n", encoding="utf-8")
        negative = [
            [python, "-B", str(SCRIPT), "--source-root", str(SKILL_SOURCE), "--scope", "unknown", "--output", str(temp_root / "unknown")],
            [python, "-B", str(SCRIPT), "--source-root", str(temp_root / "missing"), "--scope", "local-three-line", "--output", str(temp_root / "missing-out")],
            [python, "-B", str(SCRIPT), "--source-root", str(SKILL_SOURCE), "--scope", "local-three-line", "--output", str(nonempty)],
        ]
        for command in negative:
            result = subprocess.run(command, text=True, capture_output=True)
            commands.append({"argv": command, "returncode": result.returncode, "stdout": result.stdout, "stderr": result.stderr})
            assert result.returncode != 0
        assert (nonempty / "sentinel.txt").read_text(encoding="utf-8") == "preserve\n"
        local_text = "\n".join(path.read_text(encoding="utf-8") for path in (local_one / "bundle").rglob("*") if path.is_file())
        assert "checklist.md" not in local_text and "F01" not in local_text and "F20" not in local_text
        assert (full / "bundle" / "references" / "checklist.md").is_file()
        return {
            "deterministic_local_manifest": True,
            "local_excludes_checklist_and_ids": True,
            "full_includes_checklist": True,
            "negative_cases_refused": 3,
            "nonempty_sentinel_preserved": True,
            "commands": commands,
        }


def docx_result(path: Path) -> dict:
    before = Document(SOURCE)
    after = Document(path)
    assert [p.text for p in before.paragraphs] == [p.text for p in after.paragraphs]
    assert [[[c.text for c in row.cells] for row in table.rows] for table in before.tables] == [
        [[c.text for c in row.cells] for row in table.rows] for table in after.tables
    ]
    assert [(len(table.rows), len(table.columns)) for table in before.tables] == [(len(table.rows), len(table.columns)) for table in after.tables]
    with ZipFile(SOURCE) as old_zip, ZipFile(path) as new_zip:
        assert set(old_zip.namelist()) == set(new_zip.namelist())
        changed = sorted(name for name in old_zip.namelist() if hashlib.sha256(old_zip.read(name)).digest() != hashlib.sha256(new_zip.read(name)).digest())
        old_root = etree.fromstring(old_zip.read("word/document.xml"))
        new_root = etree.fromstring(new_zip.read("word/document.xml"))
    assert changed == ["word/document.xml"]
    old_body = old_root.find(W + "body")
    new_body = new_root.find(W + "body")
    old_outside = [etree.tostring(node, method="c14n") for node in old_body if node.tag != W + "tbl"]
    new_outside = [etree.tostring(node, method="c14n") for node in new_body if node.tag != W + "tbl"]
    assert old_outside == new_outside
    table = new_body.find(W + "tbl")
    rows = table.findall(W + "tr")

    def edge(cell, name):
        element = cell.find(W + "tcPr/" + W + "tcBorders/" + W + name)
        return None if element is None else (element.get(W + "val"), element.get(W + "sz"))

    for cell in rows[0].findall(W + "tc"):
        assert edge(cell, "top") == ("single", "8")
        assert edge(cell, "bottom") == ("single", "4")
    for cell in rows[-1].findall(W + "tc"):
        assert edge(cell, "bottom") == ("single", "8")
    for row in rows:
        for cell in row.findall(W + "tc"):
            for name in ("left", "right", "insideH", "insideV", "tl2br", "tr2bl"):
                assert edge(cell, name)[0] == "nil"
            shading = cell.find(W + "tcPr/" + W + "shd")
            assert shading is not None and shading.get(W + "val") == "clear"
    return {
        "sha256": digest(path),
        "changed_parts": changed,
        "outside_table_xml_equal": True,
        "text_and_shape_equal": True,
        "three_line_ooxml": True,
    }


def main() -> None:
    p01_full = ROOT / "runs" / "p01-full-snapshot" / "formatted.docx"
    p01_local = ROOT / "runs" / "p01-local-bundle-v2" / "formatted.docx"
    p02_dir = ROOT / "runs" / "p02-boundary"
    p04_dir = ROOT / "runs" / "p04-not-applicable"
    audit = ROOT / "runs" / "p03-full-audit" / "audit.md"
    audit_text = audit.read_text(encoding="utf-8")
    ids = [f"F{index:02d}" for index in range(1, 21)]
    for item in ids:
        assert audit_text.count(f"| {item} |") == 1
    assert not (p02_dir / "formatted.docx").exists()
    assert not any(path.suffix == ".docx" for path in p04_dir.iterdir())
    output = {
        "source_sha256": digest(SOURCE),
        "boundary_source_sha256": digest(BOUNDARY),
        "manifests": [verify_manifest(LOCAL_BUNDLE, "local-three-line"), verify_manifest(FULL_BUNDLE, "full-audit")],
        "staging_script_import_audit": verify_staging_imports(),
        "staging_tests": run_stage_tests(),
        "p01_full_snapshot": docx_result(p01_full),
        "p01_local_bundle": docx_result(p01_local),
        "p02_no_formatted_output": True,
        "p03_all_twenty_ids_once": True,
        "p03_checklist_read_logged": "references/checklist.md" in (ROOT / "runs" / "p03-full-audit" / "execution-log.md").read_text(encoding="utf-8"),
        "p04_no_docx_output": True,
    }
    (ROOT / "parent-verification.json").write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(ROOT / "parent-verification.json")


if __name__ == "__main__":
    main()
