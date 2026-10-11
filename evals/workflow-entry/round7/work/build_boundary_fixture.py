from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUTPUT = Path(__file__).resolve().parents[1] / "materials" / "boundary-crossing-merge.docx"


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

title = doc.add_paragraph("Synthetic respiratory outcome table")
title.style = doc.styles["Title"]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph(
    "This public synthetic fixture tests whether a table-formatting workflow "
    "stops safely when a merged cell crosses the declared header boundary."
)

table = doc.add_table(rows=5, cols=3)
table.style = "Table Grid"
values = [
    ["Treatment", "Day 0", "Day 28"],
    ["Treatment", "Mean (SD)", "Mean (SD)"],
    ["Treatment", "68.0 (7.4)", "72.1 (6.9)"],
    ["Comparator", "67.5 (7.2)", "69.0 (7.0)"],
    ["Control", "66.9 (7.7)", "67.1 (7.5)"],
]
for row_index, row_values in enumerate(values):
    for column_index, value in enumerate(row_values):
        paragraph = table.cell(row_index, column_index).paragraphs[0]
        paragraph.add_run(value)
        paragraph.paragraph_format.space_after = Pt(0)

# Declared header_rows=2 would put this merge across the header/body boundary.
table.cell(1, 0).merge(table.cell(2, 0))

doc.add_paragraph("Values are synthetic. SD, standard deviation.")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
